import base64
import hashlib
import io
import json
import os
import secrets
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Optional

import psycopg2
from cryptography.fernet import Fernet

from access import current_user
import sbis_jobs as jobs
import sbis_client as sbis

SCHEMA = os.environ.get('MAIN_DB_SCHEMA', 'public')
CORS = {
    'Access-Control-Allow-Origin': '*',
    'Access-Control-Allow-Methods': 'GET, POST, OPTIONS',
    'Access-Control-Allow-Headers': 'Content-Type, X-Authorization',
    'Access-Control-Max-Age': '86400',
}
DEFAULT_EXCLUDE_SBER = ['Инкассация']
_CTX: Dict[str, Any] = {}


def _resp(status: int, payload: Any) -> Dict[str, Any]:
    return {'statusCode': status, 'headers': {**CORS, 'Content-Type': 'application/json'},
            'body': json.dumps(payload, ensure_ascii=False, default=str)}


def _err(text: str, status: int = 400) -> Dict[str, Any]:
    return _resp(status, {'error': text})


def _fernet() -> Fernet:
    seed = (os.environ.get('SBER_VAULT_KEY') or os.environ.get('AWS_SECRET_ACCESS_KEY') or '').strip()
    if not seed:
        raise sbis.SbisError('Не настроено шифрование хранилища')
    return Fernet(base64.urlsafe_b64encode(hashlib.sha256(('sber-vault-v2:' + seed).encode()).digest()))


def _vget(cur, key: str) -> str:
    cur.execute(f'SELECT value FROM {SCHEMA}.sber_vault WHERE key = %s', ('sbis:' + key,))
    r = cur.fetchone()
    if not r:
        return ''
    try:
        return _fernet().decrypt(r[0].encode()).decode()
    except Exception:
        return ''


def _vset(cur, key: str, value: str) -> None:
    cur.execute(f'INSERT INTO {SCHEMA}.sber_vault (key, value, updated_at) VALUES (%s, %s, now()) '
                f'ON CONFLICT (key) DO UPDATE SET value = EXCLUDED.value, updated_at = now()',
                ('sbis:' + key, _fernet().encrypt(value.encode()).decode()))


def _settings(cur) -> Dict[str, Any]:
    s = jobs.jget(cur, SCHEMA, 'sbis')
    s.setdefault('points', [])
    s.setdefault('point_names', {})
    s['include'] = {**jobs.DEFAULT_INCLUDE, **(s.get('include') or {})}
    s.setdefault('exclude_sber', DEFAULT_EXCLUDE_SBER)
    return s


def _client(cur, fresh: bool = False) -> sbis.Client:
    cid, sec, key = _vget(cur, 'client_id'), _vget(cur, 'app_secret'), _vget(cur, 'secret_key')
    if not (cid and sec and key):
        raise sbis.SbisError('Ключи СБИС не заполнены — откройте Админ-панель → СБИС Престо')
    tok = '' if fresh else _vget(cur, 'token')
    if not tok:
        tok = sbis.login(cid, sec, key)
        _vset(cur, 'token', tok)
    return sbis.Client(tok)


def _with_client(cur, fn):
    try:
        return fn(_client(cur))
    except sbis.SbisError as e:
        if str(e) != 'auth':
            raise
    try:
        return fn(_client(cur, fresh=True))
    except sbis.SbisError as e:
        if str(e) == 'auth':
            raise sbis.SbisError('СБИС отклонил доступ: проверьте ключи и права приложения на «Розницу/Presto»')
        raise


def _budget() -> float:
    ctx = _CTX.get('ctx')
    left = None
    try:
        left = ctx.get_remaining_time_in_millis() / 1000.0 if ctx and hasattr(ctx, 'get_remaining_time_in_millis') else None
    except Exception:
        left = None
    if left is None:
        left = 4.5
    return max(0.5, min(25.0, left - 2.5))


def _work(conn, cur) -> Dict[str, Any]:
    s = _settings(cur)
    return _with_client(cur, lambda c: jobs.run(conn, cur, SCHEMA, c, s, _budget()))


def _status(cur, deputy: bool) -> Dict[str, Any]:
    s = _settings(cur)
    has_keys = bool(_vget(cur, 'client_id') and _vget(cur, 'app_secret') and _vget(cur, 'secret_key'))
    cur.execute(f'SELECT min(day), max(day), count(*) FROM {SCHEMA}.sbis_daily')
    mn, mx, n = cur.fetchone()
    return {'configured': has_keys and bool(s['points']), 'has_keys': has_keys, 'points': s['points'],
            'point_names': s['point_names'], 'include': s['include'], 'exclude_sber': s['exclude_sber'],
            'backfill_done': s.get('backfill_done'), 'last_auto': s.get('last_auto'),
            'client_id': (_vget(cur, 'client_id')[:4] + '…') if has_keys and deputy else None,
            'cron_key': s.get('cron_key') if deputy else None,
            'data_from': mn, 'data_to': mx, 'days': n, 'job': jobs.progress(jobs.jget(cur, SCHEMA, 'sbis_job'))}


def _report_data(cur, d_from: str, d_to: str) -> Dict[str, Any]:
    s = _settings(cur)
    inc = s['include']
    pts = [int(p) for p in s['points']]
    cur.execute(f'SELECT day, SUM(cash), SUM(bank), SUM(certificate), SUM(salary), SUM(returns), SUM(checks) '
                f'FROM {SCHEMA}.sbis_daily WHERE point_id = ANY(%s) AND day BETWEEN %s AND %s GROUP BY 1', (pts or [0], d_from, d_to))
    sb = {r[0].isoformat(): {'cash': float(r[1] or 0), 'bank': float(r[2] or 0), 'certificate': float(r[3] or 0),
                             'salary': float(r[4] or 0), 'returns': float(r[5] or 0), 'checks': int(r[6] or 0)} for r in cur.fetchall()}
    for v in sb.values():
        v['turnover'] = round(sum(v[k] for k in jobs.KINDS if inc.get(k)), 2)
    excl = [x.lower() for x in s['exclude_sber']]
    cur.execute(f"SELECT e.entry_date, COALESCE(c.name, 'Без статьи'), SUM(e.amount) FROM {SCHEMA}.acc_entries e "
                f"LEFT JOIN {SCHEMA}.acc_categories c ON c.id = e.category_id WHERE e.source = 'bank' AND e.kind = 'income' "
                f'AND NOT e.void AND e.entry_date BETWEEN %s AND %s GROUP BY 1, 2', (d_from, d_to))
    bank: Dict[str, Dict[str, float]] = {}
    bank_cats: Dict[str, float] = {}
    for d, cat, total in cur.fetchall():
        b = bank.setdefault(d.isoformat(), {'in_total': 0.0, 'excluded': 0.0})
        b['in_total'] += float(total or 0)
        bank_cats[cat] = bank_cats.get(cat, 0.0) + float(total or 0)
        if cat.lower() in excl:
            b['excluded'] += float(total or 0)
    days = []
    d, end = date.fromisoformat(d_from), date.fromisoformat(d_to)
    while d <= end:
        k = d.isoformat()
        a = sb.get(k, {})
        b = bank.get(k, {'in_total': 0.0, 'excluded': 0.0})
        sber_in = round(b['in_total'] - b['excluded'], 2)
        if a or b['in_total']:
            days.append({'date': k, 'cash': a.get('cash', 0.0), 'certificate': a.get('certificate', 0.0), 'salary': a.get('salary', 0.0),
                         'bank_sbis': a.get('bank', 0.0), 'returns': a.get('returns', 0.0), 'checks': a.get('checks', 0),
                         'turnover': a.get('turnover', 0.0), 'sber': sber_in, 'sber_excluded': round(b['excluded'], 2),
                         'total': round(a.get('turnover', 0.0) + sber_in, 2)})
        d += timedelta(days=1)
    tot = {k: round(sum(x[k] for x in days), 2) for k in ('cash', 'certificate', 'salary', 'bank_sbis', 'returns', 'turnover', 'sber', 'sber_excluded', 'total')}
    tot['checks'] = sum(x['checks'] for x in days)
    return {'from': d_from, 'to': d_to, 'days': days, 'totals': tot, 'include': inc, 'exclude_sber': s['exclude_sber'],
            'bank_categories': [{'category': k, 'total': round(v, 2), 'excluded': k.lower() in excl} for k, v in sorted(bank_cats.items(), key=lambda x: -x[1])]}


def _xlsx(cur, d_from: str, d_to: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    rep = _report_data(cur, d_from, d_to)
    wb = Workbook()
    bold, head = Font(bold=True), PatternFill('solid', fgColor='DCEFE3')
    money = '#,##0.00'

    def table(ws, start_row: int, headers: List[str], rows: List[List[Any]], widths: List[int], money_from: int = 1, total: Optional[List[Any]] = None):
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=start_row, column=i, value=h)
            c.font, c.fill = bold, head
            c.alignment = Alignment(wrap_text=True, vertical='center')
        r = start_row
        for row in rows + ([total] if total else []):
            r += 1
            for i, v in enumerate(row, 1):
                c = ws.cell(row=r, column=i, value=v)
                if i > money_from and isinstance(v, (int, float)):
                    c.number_format = money
                if total and row is total:
                    c.font = bold
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)

    fmt_d = lambda s: f'{s[8:10]}.{s[5:7]}.{s[:4]}'
    t = rep['totals']
    ws = wb.active
    ws.title = 'Итог'
    ws['A1'] = f'Выручка СБИС Престо + СберБизнес за {fmt_d(d_from)} — {fmt_d(d_to)}'
    ws['A1'].font = Font(bold=True, size=14)
    names = {'cash': 'наличные', 'certificate': 'сертификаты', 'salary': 'под зарплату', 'bank': 'безнал/QR/терминал'}
    ws['A2'] = 'В кассовый оборот СБИС входит: ' + ', '.join(n for k, n in names.items() if rep['include'].get(k)) + \
               '. Не входит: ' + (', '.join(n for k, n in names.items() if not rep['include'].get(k)) or '—') + ' (приходит на расчётный счёт).'
    ws['A3'] = 'Из поступлений Сбера не учтены статьи: ' + (', '.join(rep['exclude_sber']) or '—')
    summary = [
        ['СБИС Престо — кассовый оборот', t['turnover']],
        ['   в т.ч. наличные', t['cash']],
        ['   в т.ч. сертификаты', t['certificate']],
        ['   в т.ч. под зарплату', t['salary']],
        ['СБИС — безнал (не в итоге, идёт на счёт)', t['bank_sbis']],
        ['СБИС — возвраты (уже вычтены)', t['returns']],
        ['СБИС — чеков', t['checks']],
        ['СберБизнес — поступления на счёт', t['sber']],
        ['СберБизнес — не учтено (' + ', '.join(rep['exclude_sber']) + ')', t['sber_excluded']],
    ]
    table(ws, 5, ['Показатель', 'Сумма, ₽'], summary, [52, 20], total=['ИТОГ ОБЩИЙ (СБИС + Сбер)', t['total']])
    ws.freeze_panes = None

    ws2 = wb.create_sheet('По дням')
    rows = [[fmt_d(x['date']), x['cash'], x['certificate'], x['salary'], x['turnover'], x['bank_sbis'], x['returns'], x['checks'], x['sber'], x['total']]
            for x in rep['days']]
    table(ws2, 1, ['Дата', 'СБИС наличные', 'СБИС сертификаты', 'СБИС под з/п', 'СБИС кассовый оборот', 'СБИС безнал (не в итоге)',
                   'СБИС возвраты', 'Чеков', 'Сбер поступления', 'ИТОГ ОБЩИЙ'], rows, [12, 15, 15, 14, 18, 18, 14, 9, 18, 18],
          total=['Итого', t['cash'], t['certificate'], t['salary'], t['turnover'], t['bank_sbis'], t['returns'], t['checks'], t['sber'], t['total']])

    months: Dict[str, Dict[str, float]] = {}
    for x in rep['days']:
        m = months.setdefault(x['date'][:7], {'turnover': 0.0, 'sber': 0.0, 'total': 0.0})
        for k in m:
            m[k] += x[k]
    ws3 = wb.create_sheet('По месяцам')
    table(ws3, 1, ['Месяц', 'СБИС кассовый оборот', 'Сбер поступления', 'ИТОГ ОБЩИЙ'],
          [[f'{k[5:7]}.{k[:4]}', round(v['turnover'], 2), round(v['sber'], 2), round(v['total'], 2)] for k, v in sorted(months.items())],
          [12, 22, 20, 20], total=['Итого', t['turnover'], t['sber'], t['total']])

    ws4 = wb.create_sheet('Сбер по статьям')
    table(ws4, 1, ['Статья поступления', 'Сумма, ₽', 'В итоге'],
          [[c['category'], c['total'], 'нет' if c['excluded'] else 'да'] for c in rep['bank_categories']], [40, 18, 10])

    cur.execute(f"SELECT e.entry_date, CASE WHEN e.kind = 'income' THEN 'Доход' ELSE 'Расход' END, COALESCE(c.name, 'Без статьи'), "
                f"e.amount, COALESCE(cp.name, ''), CASE e.source WHEN 'bank' THEN 'СберБизнес' WHEN 'sbis' THEN 'СБИС Престо' "
                f"WHEN 'collection' THEN 'Инкассация' WHEN 'salary' THEN 'Зарплата' WHEN 'salary_week' THEN 'Зарплата' ELSE 'Вручную' END, "
                f"COALESCE(e.comment, '') FROM {SCHEMA}.acc_entries e LEFT JOIN {SCHEMA}.acc_categories c ON c.id = e.category_id "
                f'LEFT JOIN {SCHEMA}.acc_counterparties cp ON cp.id = e.counterparty_id '
                f'WHERE NOT e.void AND e.entry_date BETWEEN %s AND %s ORDER BY e.entry_date, e.id', (d_from, d_to))
    ops = [[r[0].strftime('%d.%m.%Y'), r[1], r[2], float(r[3]), r[4], r[5], r[6][:500]] for r in cur.fetchall()]
    ws5 = wb.create_sheet('Все операции')
    inc_sum = round(sum(o[3] for o in ops if o[1] == 'Доход'), 2)
    exp_sum = round(sum(o[3] for o in ops if o[1] == 'Расход'), 2)
    table(ws5, 1, ['Дата', 'Тип', 'Статья', 'Сумма, ₽', 'Контрагент', 'Источник', 'Комментарий'], ops, [12, 9, 32, 15, 30, 14, 70], money_from=3)
    r = len(ops) + 3
    for label, v in (('Доходы', inc_sum), ('Расходы', exp_sum), ('Разница', round(inc_sum - exp_sum, 2))):
        ws5.cell(row=r, column=3, value=label).font = bold
        c = ws5.cell(row=r, column=4, value=v)
        c.font, c.number_format = bold, money
        r += 1

    cats: Dict[tuple, float] = {}
    for o in ops:
        cats[(o[1], o[2])] = cats.get((o[1], o[2]), 0.0) + o[3]
    ws6 = wb.create_sheet('Статьи')
    table(ws6, 1, ['Тип', 'Статья', 'Сумма, ₽'], [[k[0], k[1], round(v, 2)] for k, v in sorted(cats.items(), key=lambda x: (x[0][0] != 'Доход', -x[1]))],
          [10, 40, 18], money_from=2)

    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


def _d(v: Any) -> Optional[str]:
    try:
        return date.fromisoformat(str(v)[:10]).isoformat()
    except ValueError:
        return None


def _auto(conn, cur) -> Dict[str, Any]:
    """Планировщик: если нет задания — ставит догрузку последних дней (раз в час), затем выполняет шаг."""
    s = _settings(cur)
    if not s['points']:
        return {'skipped': 'Не выбраны точки продаж'}
    job = jobs.jget(cur, SCHEMA, 'sbis_job')
    now = datetime.now()
    last = s.get('last_auto') or ''
    if (not job or job.get('done')) and last[:13] != now.strftime('%Y-%m-%dT%H'):
        d_to = date.today()
        jobs.start(cur, SCHEMA, 'auto', (d_to - timedelta(days=2)).isoformat(), d_to.isoformat(), s['points'], 'Планировщик')
        s = jobs.jget(cur, SCHEMA, 'sbis')
        s['last_auto'] = now.strftime('%Y-%m-%dT%H:%M')
        jobs.jset(cur, SCHEMA, 'sbis', s, 'Планировщик')
    return _work(conn, cur)


def handle(event: Dict[str, Any], context: Any) -> Dict[str, Any]:
    '''СБИС Престо → бухгалтерия: выручка по дням (статья «Кассовый оборот», без безнала, который приходит на счёт),
    разовая выгрузка за 2 года, догрузка по планировщику, общий итог СБИС + СберБизнес и отчёт xlsx за промежуток.'''
    method = event.get('httpMethod')
    if method == 'OPTIONS':
        return {'statusCode': 200, 'headers': CORS, 'body': ''}
    _CTX['ctx'] = context
    qs = event.get('queryStringParameters') or {}
    conn = psycopg2.connect(os.environ['DATABASE_URL'])
    conn.autocommit = True
    cur = conn.cursor()
    try:
        if method == 'GET' and qs.get('cron') is not None:
            key = _settings(cur).get('cron_key')
            if not key or qs.get('cron') != key:
                return _err('Нет доступа', 403)
            try:
                return _resp(200, {'ok': True, **_auto(conn, cur)})
            except sbis.SbisError as e:
                return _resp(200, {'ok': False, 'error': str(e)})
        if method != 'POST':
            return _err('Метод не поддерживается', 405)
        body = json.loads(event.get('body') or '{}')
        user = current_user(cur, SCHEMA, event)
        if not user:
            return _err('Нужно войти в систему', 401)
        action = (body.get('action') or '')[5:]
        deputy = user['is_deputy']
        can = user['actions']
        try:
            return _action(conn, cur, user, deputy, can, action, body)
        except sbis.SbisError as e:
            return _err(str(e), 502)
    finally:
        cur.close()
        conn.close()


def _action(conn, cur, user, deputy: bool, can, action: str, body: Dict[str, Any]) -> Dict[str, Any]:
    in_acc = '/accounting' in user['sections']
    if action == 'status':
        if not (deputy or in_acc):
            return _err('Нет доступа', 403)
        return _resp(200, _status(cur, deputy))

    if action in ('report', 'xlsx'):
        if not (deputy or (in_acc and 'accounting.view_reports' in can)):
            return _err('Нет права смотреть отчёты', 403)
        d_from, d_to = _d(body.get('date_from')), _d(body.get('date_to'))
        if not d_from or not d_to or d_from > d_to:
            return _err('Укажите промежуток дат')
        if action == 'report':
            return _resp(200, _report_data(cur, d_from, d_to))
        return _resp(200, {'file': _xlsx(cur, d_from, d_to), 'name': f'Выручка_СБИС_Сбер_{d_from}_{d_to}.xlsx'})

    if action in ('load', 'continue'):
        if not (deputy or (in_acc and 'accounting.edit_entries' in can)):
            return _err('Нет права загружать данные', 403)
        s = _settings(cur)
        if not s['points']:
            return _err('Сначала настройте СБИС в админ-панели и выберите точки продаж')
        if action == 'load':
            d_from, d_to = _d(body.get('date_from')), _d(body.get('date_to'))
            if not d_from or not d_to or d_from > d_to:
                return _err('Укажите промежуток дат')
            if d_to > date.today().isoformat():
                d_to = date.today().isoformat()
            if (date.fromisoformat(d_to) - date.fromisoformat(d_from)).days > 400:
                return _err('Не больше 13 месяцев за раз (за 2 года — разовая выгрузка в админ-панели)')
            msg = jobs.start(cur, SCHEMA, 'range', d_from, d_to, s['points'], user['name'])
            if msg:
                return _err(msg, 409)
        return _resp(200, {'job': _work(conn, cur)})

    if not deputy:
        return _err('Настройки СБИС меняет только заместитель директора', 403)

    if action == 'keys_save':
        vals = {k: str(body.get(k) or '').strip() for k in ('client_id', 'app_secret', 'secret_key')}
        if not all(vals.values()):
            return _err('Заполните все три поля из карточки приложения СБИС')
        sbis.login(vals['client_id'], vals['app_secret'], vals['secret_key'])
        for k, v in vals.items():
            _vset(cur, k, v)
        _vset(cur, 'token', '')
        s = jobs.jget(cur, SCHEMA, 'sbis')
        s.setdefault('cron_key', secrets.token_hex(16))
        jobs.jset(cur, SCHEMA, 'sbis', s, user['name'])
        return _resp(200, {'ok': True, 'status': _status(cur, True)})
    if action == 'keys_delete':
        cur.execute(f"DELETE FROM {SCHEMA}.sber_vault WHERE key LIKE 'sbis:%%'")
        return _resp(200, {'ok': True, 'status': _status(cur, True)})
    if action == 'points':
        return _resp(200, {'points': _with_client(cur, lambda c: c.points())})
    if action == 'settings_save':
        s = jobs.jget(cur, SCHEMA, 'sbis')
        pts = body.get('points') or []
        s['points'] = [int(p) for p in pts if str(p).isdigit()]
        s['point_names'] = {str(k): str(v)[:120] for k, v in (body.get('point_names') or {}).items()}
        s['include'] = {k: bool((body.get('include') or {}).get(k, jobs.DEFAULT_INCLUDE[k])) for k in jobs.KINDS}
        s['exclude_sber'] = [str(x).strip()[:100] for x in (body.get('exclude_sber') or []) if str(x).strip()]
        s.setdefault('cron_key', secrets.token_hex(16))
        jobs.jset(cur, SCHEMA, 'sbis', s, user['name'])
        return _resp(200, {'ok': True, 'status': _status(cur, True)})
    if action == 'backfill':
        s = _settings(cur)
        if not s['points']:
            return _err('Сначала выберите точки продаж и сохраните')
        if s.get('backfill_done') and not body.get('force'):
            return _err('Разовая выгрузка за 2 года уже выполнена ' + s['backfill_done'], 409)
        job = jobs.jget(cur, SCHEMA, 'sbis_job')
        if job and job.get('kind') == 'backfill' and not job.get('done'):
            return _resp(200, {'job': _work(conn, cur)})
        today = date.today()
        jobs.start(cur, SCHEMA, 'backfill', today.replace(year=today.year - 2).isoformat(), today.isoformat(), s['points'], user['name'])
        return _resp(200, {'job': _work(conn, cur)})
    if action == 'cron_reset':
        s = jobs.jget(cur, SCHEMA, 'sbis')
        s['cron_key'] = secrets.token_hex(16)
        jobs.jset(cur, SCHEMA, 'sbis', s, user['name'])
        return _resp(200, {'ok': True, 'status': _status(cur, True)})
    return _err('Неизвестное действие')


def finders(qs):
    if qs.get('cron') is not None:
        return [('app_settings', "key = 'sbis' AND value->>'cron_key' = %s", [qs.get('cron') or ''])]
    return []
